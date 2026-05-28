#!/usr/bin/env python3
"""Import CSV rows into an existing ODS or XLSX spreadsheet."""

from __future__ import annotations

import argparse
import csv
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence
from zipfile import BadZipFile


SUPPORTED_SPREADSHEET_EXTENSIONS = {".ods", ".xlsx"}


class CsvToCalcError(Exception):
    """User-facing error with a CLI exit code."""

    def __init__(self, message: str, exit_code: int = 1) -> None:
        super().__init__(message)
        self.exit_code = exit_code


@dataclass(frozen=True)
class ImportConfig:
    csv_path: Path
    target_path: Path
    sheet_name: str | None
    start_row: int
    start_col: int
    delimiter: str
    encoding: str
    has_header: bool
    dry_run: bool


def positive_int(value: str) -> int:
    try:
        parsed = int(value)
    except ValueError as exc:
        raise argparse.ArgumentTypeError("musi być liczbą całkowitą większą od 0") from exc

    if parsed <= 0:
        raise argparse.ArgumentTypeError("musi być większe od 0")
    return parsed


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Importuje dane z pliku CSV do istniejącego pliku .ods lub .xlsx.",
    )
    parser.add_argument("csv_file", help="ścieżka do źródłowego pliku .csv")
    parser.add_argument("target_file", help="ścieżka do istniejącego pliku .ods lub .xlsx")
    parser.add_argument("--sheet", help="nazwa arkusza docelowego; domyślnie aktywny/pierwszy arkusz")
    parser.add_argument("--start-row", type=positive_int, default=1, help="pierwszy wiersz zapisu, numerowany od 1")
    parser.add_argument("--start-col", type=positive_int, default=1, help="pierwsza kolumna zapisu, numerowana od 1")
    parser.add_argument("--delimiter", default=",", help="separator CSV; domyślnie przecinek")
    parser.add_argument("--encoding", default="utf-8", help="kodowanie CSV; domyślnie utf-8")
    parser.add_argument(
        "--has-header",
        "--skip-header",
        action="store_true",
        dest="has_header",
        help="traktuje pierwszy wiersz CSV jako nagłówek i nie zapisuje go do arkusza",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="sprawdza dane i pokazuje plan importu bez zapisywania pliku",
    )
    return parser


def validate_args(args: argparse.Namespace) -> ImportConfig:
    csv_path = Path(args.csv_file)
    target_path = Path(args.target_file)

    if len(args.delimiter) != 1:
        raise CsvToCalcError("Separator CSV musi być pojedynczym znakiem.", exit_code=2)
    if csv_path.suffix.lower() != ".csv":
        raise CsvToCalcError("Plik źródłowy musi mieć rozszerzenie .csv.", exit_code=2)
    if target_path.suffix.lower() not in SUPPORTED_SPREADSHEET_EXTENSIONS:
        raise CsvToCalcError("Plik docelowy musi mieć rozszerzenie .ods albo .xlsx.", exit_code=2)
    if not csv_path.exists():
        raise CsvToCalcError(f"Nie znaleziono pliku CSV: {csv_path}", exit_code=2)
    if not csv_path.is_file():
        raise CsvToCalcError(f"Ścieżka CSV nie jest plikiem: {csv_path}", exit_code=2)
    if not target_path.exists():
        raise CsvToCalcError(f"Nie znaleziono pliku docelowego: {target_path}", exit_code=2)
    if not target_path.is_file():
        raise CsvToCalcError(f"Ścieżka docelowa nie jest plikiem: {target_path}", exit_code=2)

    return ImportConfig(
        csv_path=csv_path,
        target_path=target_path,
        sheet_name=args.sheet,
        start_row=args.start_row,
        start_col=args.start_col,
        delimiter=args.delimiter,
        encoding=args.encoding,
        has_header=args.has_header,
        dry_run=args.dry_run,
    )


def read_csv(path: Path | str, delimiter: str = ",", encoding: str = "utf-8") -> list[list[str]]:
    try:
        with Path(path).open("r", newline="", encoding=encoding) as csv_file:
            return list(csv.reader(csv_file, delimiter=delimiter))
    except FileNotFoundError as exc:
        raise CsvToCalcError(f"Nie znaleziono pliku CSV: {path}") from exc
    except PermissionError as exc:
        raise CsvToCalcError(f"Brak uprawnień do odczytu pliku CSV: {path}") from exc
    except UnicodeDecodeError as exc:
        raise CsvToCalcError(f"Nie można odczytać pliku CSV w kodowaniu {encoding}: {path}") from exc
    except csv.Error as exc:
        raise CsvToCalcError(f"Nie można przetworzyć pliku CSV {path}: {exc}") from exc
    except OSError as exc:
        raise CsvToCalcError(f"Błąd odczytu pliku CSV {path}: {exc}") from exc


def rows_to_write(rows: Sequence[Sequence[str]], has_header: bool) -> list[Sequence[str]]:
    if has_header and rows:
        return list(rows[1:])
    return list(rows)


def count_rows_to_write(rows: Sequence[Sequence[str]], has_header: bool) -> int:
    return len(rows_to_write(rows, has_header))


def import_rows(config: ImportConfig, rows: Sequence[Sequence[str]]) -> str:
    extension = config.target_path.suffix.lower()
    if extension == ".xlsx":
        return import_rows_to_xlsx(config, rows)
    if extension == ".ods":
        return import_rows_to_ods(config, rows)
    raise CsvToCalcError("Nieobsługiwany format pliku docelowego.")


def import_rows_to_xlsx(config: ImportConfig, rows: Sequence[Sequence[str]]) -> str:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise CsvToCalcError("Brak biblioteki openpyxl. Zainstaluj zależności z requirements.txt.") from exc

    try:
        workbook = load_workbook(config.target_path)
    except (BadZipFile, InvalidFileException, KeyError) as exc:
        raise CsvToCalcError(f"Plik .xlsx ma nieprawidłowy format: {config.target_path}") from exc
    except OSError as exc:
        raise CsvToCalcError(f"Błąd otwierania pliku .xlsx {config.target_path}: {exc}") from exc

    if config.sheet_name:
        if config.sheet_name not in workbook.sheetnames:
            raise CsvToCalcError(f"Nie znaleziono arkusza '{config.sheet_name}' w pliku {config.target_path}.")
        worksheet = workbook[config.sheet_name]
    else:
        worksheet = workbook.active

    for row_offset, row in enumerate(rows):
        for col_offset, value in enumerate(row):
            worksheet.cell(
                row=config.start_row + row_offset,
                column=config.start_col + col_offset,
                value=value,
            )

    try:
        workbook.save(config.target_path)
    except PermissionError as exc:
        raise CsvToCalcError(f"Brak uprawnień do zapisu pliku .xlsx: {config.target_path}") from exc
    except OSError as exc:
        raise CsvToCalcError(f"Błąd zapisu pliku .xlsx {config.target_path}: {exc}") from exc

    return worksheet.title


def import_rows_to_ods(config: ImportConfig, rows: Sequence[Sequence[str]]) -> str:
    try:
        from odf.opendocument import load
        from odf.table import Table, TableCell, TableRow
        from odf.text import P
    except ImportError as exc:
        raise CsvToCalcError("Brak biblioteki odfpy. Zainstaluj zależności z requirements.txt.") from exc

    try:
        document = load(str(config.target_path))
    except OSError as exc:
        raise CsvToCalcError(f"Błąd otwierania pliku .ods {config.target_path}: {exc}") from exc
    except Exception as exc:
        raise CsvToCalcError(f"Plik .ods ma nieprawidłowy format: {config.target_path}") from exc

    tables = document.spreadsheet.getElementsByType(Table)
    if config.sheet_name:
        table = next((item for item in tables if item.getAttribute("name") == config.sheet_name), None)
        if table is None:
            raise CsvToCalcError(f"Nie znaleziono arkusza '{config.sheet_name}' w pliku {config.target_path}.")
    elif tables:
        table = tables[0]
    else:
        raise CsvToCalcError(f"Plik .ods nie zawiera arkuszy: {config.target_path}")

    for row_offset, row in enumerate(rows):
        table_row = _ensure_ods_row(table, config.start_row + row_offset - 1, TableRow)
        for col_offset, value in enumerate(row):
            table_cell = _ensure_ods_cell(table_row, config.start_col + col_offset - 1, TableCell)
            _set_ods_cell_text(table_cell, value, P)

    try:
        document.save(str(config.target_path))
    except PermissionError as exc:
        raise CsvToCalcError(f"Brak uprawnień do zapisu pliku .ods: {config.target_path}") from exc
    except OSError as exc:
        raise CsvToCalcError(f"Błąd zapisu pliku .ods {config.target_path}: {exc}") from exc

    return table.getAttribute("name") or "pierwszy arkusz"


def _ensure_ods_row(table: object, index: int, row_type: type) -> object:
    rows = table.getElementsByType(row_type)
    while len(rows) <= index:
        row = row_type()
        table.addElement(row)
        rows.append(row)
    return rows[index]


def _ensure_ods_cell(row: object, index: int, cell_type: type) -> object:
    cells = row.getElementsByType(cell_type)
    while len(cells) <= index:
        cell = cell_type()
        row.addElement(cell)
        cells.append(cell)
    return cells[index]


def _set_ods_cell_text(cell: object, value: str, paragraph_type: type) -> None:
    while cell.childNodes:
        cell.removeChild(cell.childNodes[0])
    cell.setAttribute("valuetype", "string")
    cell.addElement(paragraph_type(text=value))


def main(argv: Sequence[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    try:
        config = validate_args(args)
        csv_rows = read_csv(config.csv_path, delimiter=config.delimiter, encoding=config.encoding)
        data_rows = rows_to_write(csv_rows, config.has_header)
        planned_sheet = config.sheet_name or "aktywny/pierwszy arkusz"

        if config.dry_run:
            print(
                "Tryb dry-run: zapis pominięty. "
                f"Wiersze do zapisania: {len(data_rows)}. "
                f"Arkusz docelowy: {planned_sheet}."
            )
            return 0

        actual_sheet = import_rows(config, data_rows)
    except CsvToCalcError as exc:
        print(f"Błąd: {exc}", file=sys.stderr)
        return exc.exit_code

    print(f"Zapisano {len(data_rows)} wierszy do arkusza: {actual_sheet}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
